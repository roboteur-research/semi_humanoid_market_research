#!/usr/bin/env python3
"""Compile all companies/*/*/specs.json into study/semi_humanoid_market_2026.xlsx."""
import json
import glob
import os
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "study", "semi_humanoid_market_2026.xlsx")

IDENTITY = [
    ("robot", "Robot"), ("robot_native", "Native name"), ("company", "Company"),
    ("company_native", "Company (native)"), ("hq_country", "HQ country"),
    ("hq_city", "HQ city"), ("region", "Region"), ("status", "Status"),
    ("first_shown", "First shown"), ("target_applications", "Target applications"),
]
SPECS = [
    ("base_type", "Base type"), ("height_mm", "Height (mm)"), ("weight_kg", "Weight (kg)"),
    ("dof_total", "DoF total"), ("dof_per_arm", "DoF per arm"), ("reach_mm", "Reach (mm)"),
    ("payload_per_arm_kg", "Payload/arm (kg)"), ("payload_total_kg", "Payload total (kg)"),
    ("torso_lift", "Torso/lift"), ("max_speed_ms", "Max speed (m/s)"),
    ("battery_kwh", "Battery (kWh)"), ("runtime_h", "Runtime (h)"),
    ("actuators", "Actuators"), ("end_effector", "End effector / hands"),
    ("media_at_flange", "Media at flange"), ("sensors_head", "Sensors: head"),
    ("sensors_wrist_hand", "Sensors: wrist/hand"), ("sensors_base", "Sensors: base"),
    ("compute", "Compute"), ("software_stack", "Software stack"),
    ("safety_compliance", "Safety & compliance"), ("price", "Price"),
    ("availability", "Availability"), ("deployment_evidence", "Deployment evidence"),
]

CONF_FILL = {
    "vendor-claimed": PatternFill("solid", fgColor="FFF2CC"),   # light amber
    "third-party": PatternFill("solid", fgColor="D9EAD3"),      # light green
    "estimated": PatternFill("solid", fgColor="F4CCCC"),        # light red
}
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
THIN = Border(*[Side(style="thin", color="CCCCCC")] * 4)

STATUS_ORDER = {"shipping": 0, "announced": 1, "prototype": 2,
                "research": 3, "discontinued": 4}


def is_annex(robot):
    """Borderline/excluded entries go to a separate sheet, not the main matrix."""
    blob = " ".join([
        str(robot.get("notes", "")),
        str((robot.get("base_type") or {}).get("value", "")
            if isinstance(robot.get("base_type"), dict) else robot.get("base_type", "")),
    ]).upper()
    return "[ANNEX]" in blob


def load_all():
    robots = []
    for path in sorted(glob.glob(os.path.join(ROOT, "companies", "*", "*", "specs.json"))):
        try:
            with open(path) as f:
                data = json.load(f)
            data["_path"] = os.path.relpath(path, ROOT)
            robots.append(data)
        except Exception as e:
            print(f"WARN: failed to load {path}: {e}", file=sys.stderr)
    robots.sort(key=lambda r: (r.get("region", ""),
                               STATUS_ORDER.get(r.get("status", ""), 9),
                               r.get("company", ""), r.get("robot", "")))
    return robots


def field_value(robot, key):
    v = robot.get(key)
    if isinstance(v, dict):
        return v.get("value"), v.get("confidence", ""), v.get("source", "")
    return v, "", ""


def write_matrix(ws, robots, keys, with_conf=True):
    headers = [h for _, h in keys]
    ws.append(headers)
    for c, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r, robot in enumerate(robots, 2):
        for c, (key, _) in enumerate(keys, 1):
            val, conf, _src = field_value(robot, key)
            cell = ws.cell(row=r, column=c,
                           value=val if val is not None else "n/a")
            cell.border = THIN
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.font = Font(size=9)
            if with_conf and conf in CONF_FILL:
                cell.fill = CONF_FILL[conf]
    ws.freeze_panes = "C2"
    for c in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 14 if c > 2 else 20


def main():
    allr = load_all()
    robots = [r for r in allr if not is_annex(r)]
    annex = [r for r in allr if is_annex(r)]
    print(f"Loaded {len(allr)} robots ({len(robots)} main, {len(annex)} borderline/excluded)")
    wb = Workbook()

    ws = wb.active
    ws.title = "Overview"
    overview_keys = IDENTITY + [("base_type", "Base type"), ("price", "Price"),
                                ("payload_per_arm_kg", "Payload/arm (kg)"),
                                ("deployment_evidence", "Deployment evidence")]
    write_matrix(ws, robots, overview_keys)

    ws2 = wb.create_sheet("Full spec matrix")
    write_matrix(ws2, robots, IDENTITY + SPECS)

    for region in sorted({r.get("region", "Other") for r in robots}):
        sub = [r for r in robots if r.get("region", "Other") == region]
        wsr = wb.create_sheet(region[:28])
        write_matrix(wsr, sub, IDENTITY + SPECS)

    ws_ax = wb.create_sheet("Borderline & excluded")
    write_matrix(ws_ax, annex, IDENTITY + [("base_type", "Base type"),
                                           ("payload_per_arm_kg", "Payload/arm (kg)"),
                                           ("price", "Price")])

    ws_src = wb.create_sheet("Sources")
    ws_src.append(["Robot", "Company", "Field", "Confidence", "Source URL"])
    for c in range(1, 6):
        cell = ws_src.cell(row=1, column=c)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for robot in robots:
        for key, label in SPECS:
            val, conf, src = field_value(robot, key)
            if src:
                ws_src.append([robot.get("robot"), robot.get("company"),
                               label, conf, src])
        for src in robot.get("sources", []):
            ws_src.append([robot.get("robot"), robot.get("company"),
                           "general", "", src])

    ws_leg = wb.create_sheet("Legend")
    ws_leg.append(["Confidence colour coding (Full spec matrix)"])
    for name, fill in CONF_FILL.items():
        ws_leg.append([name])
        ws_leg.cell(row=ws_leg.max_row, column=1).fill = fill
    ws_leg.append(["'n/a' = not disclosed / unknown"])

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
